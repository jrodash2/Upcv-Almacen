from decimal import Decimal
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from django.contrib.auth.models import Group, User
from django.template import Template, Context
from django.template.loader import get_template
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import load_workbook
from pypdf import PdfReader

from .formatos_cantidades import cantidad_formato, formato_excel_cantidad
from .templatetags.moneda_filters import quetzales
from .models import Articulo, Categoria, Serie, form1h, DetalleFactura


class CantidadFormatoTests(SimpleTestCase):
    def test_enteros_con_agrupacion_sin_decimales(self):
        for valor, esperado in [('0.00', '0'), ('1.00', '1'), ('2.00', '2'),
                                ('9.00', '9'), ('10.00', '10'), ('1250.00', '1,250'),
                                ('-1250.00', '-1,250'), ('-0.00', '0'), (None, '0'), ('', '0')]:
            with self.subTest(valor=valor):
                self.assertEqual(cantidad_formato(valor), esperado)

    def test_fracciones_no_se_redondean_ni_pierden_precision(self):
        for valor, esperado in [('1.50', '1.5'), ('1.2500', '1.25'), ('1.23456789', '1.23456789'),
                                ('1250.000125', '1,250.000125'), ('0.000001', '0.000001'),
                                ('1E+4', '10,000'), ('12345678901234567890.123456789', '12,345,678,901,234,567,890.123456789')]:
            with self.subTest(valor=valor):
                numero = Decimal(valor)
                self.assertEqual(cantidad_formato(numero), esperado)
                self.assertEqual(numero, Decimal(valor))

    def test_invalido_no_se_disfraza_de_cero(self):
        for valor in ['texto', 'NaN', 'Infinity']:
            self.assertEqual(cantidad_formato(valor), valor)

    def test_moneda_conserva_dos_decimales_y_autoescape(self):
        html = Template('{% load cantidad_filters moneda_filters %}{{ unidades|cantidad_formato }} / {{ precio|quetzales }}')
        self.assertEqual(html.render(Context({'unidades': Decimal('1250.00'), 'precio': Decimal('1250')})), '1,250 / Q1,250.00')
        self.assertEqual(quetzales(Decimal('12.1')), 'Q12.10')
        self.assertEqual(quetzales(Decimal('15000.5')), 'Q15,000.50')
        self.assertEqual(Template('{% load cantidad_filters %}{{ x|cantidad_formato }}').render(Context({'x': '<b>x</b>'})), '&lt;b&gt;x&lt;/b&gt;')

    def test_formato_excel_no_oculta_fracciones(self):
        for valor, esperado in [('1250.00', '#,##0'), ('1.50', '#,##0.#'),
                                ('1.25', '#,##0.##'), ('1.000125', '#,##0.######')]:
            self.assertEqual(formato_excel_cantidad(Decimal(valor)), esperado)

    def test_templates_de_cantidades_compilan(self):
        root = Path(__file__).parent / 'templates'
        for path in root.rglob('*.html'):
            if 'cantidad_formato' in path.read_text():
                with self.subTest(template=path.name):
                    get_template(str(path.relative_to(root)))


class ExportacionCantidadTests(TestCase):
    def setUp(self):
        self.usuario = User.objects.create_user(username='formatos')
        self.usuario.groups.add(Group.objects.get_or_create(name='Administrador')[0])
        self.client.force_login(self.usuario)
        self.articulo = Articulo.objects.create(codigo='FORM-1250', nombre='Prueba de formato',
            categoria=Categoria.objects.create(nombre='Suministro'))
        Serie.objects.create(serie='FMT', numero_inicial=1, numero_final=100)
        self.factura = form1h.objects.create(numero_factura='FMT-1', estado='confirmado')
        DetalleFactura.objects.create(form1h=self.factura, articulo=self.articulo, cantidad=1250,
            precio_unitario=Decimal('12.10'), precio_total=Decimal('15125'), id_linea=1, renglon=243)

    def test_excel_detalle_contiene_numeros_y_moneda_sin_cambios(self):
        response = self.client.get(reverse('almacen:exportar_excel', args=[self.factura.pk]))
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual((sheet['A6'].value, sheet['A6'].data_type, sheet['A6'].number_format), (1250, 'n', '#,##0'))
        self.assertEqual(sheet['E6'].value, 12.1)
        self.assertEqual(sheet['E6'].number_format, '"Q"#,##0.00')
        self.assertEqual(sheet['F6'].value, 15125)

    def test_pdf_cantidad_agrupada_y_moneda_de_dos_decimales(self):
        response = self.client.get(reverse('almacen:exportar_detalle_factura_pdf', args=[self.factura.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF'))
        texto = '\n'.join(page.extract_text() or '' for page in PdfReader(BytesIO(response.content)).pages)
        self.assertIn('1,250', texto)
        self.assertIn('Q12.10', texto)
        self.assertIn('Q15,125.00', texto)

    def test_libro_excel_no_convierte_a_texto_ni_redondea_fracciones(self):
        fila = dict(renglon=243, codigo='FMT', articulo='Artículo', unidad='Unidad',
            saldo_inicial_cantidad=Decimal('1250'), saldo_inicial_monto=Decimal('20'),
            entradas_cantidad=Decimal('1.25'), entradas_monto=Decimal('12.1'),
            salidas_cantidad=Decimal('1.2345'), salidas_monto=Decimal('2'),
            saldo_final_cantidad=Decimal('1.000125'), saldo_final_monto=Decimal('30.1'), costo_promedio=Decimal('12.1'))
        datos = {'filas': [fila], 'filtros': {'renglon_inicio': 200, 'renglon_fin': 299, 'fecha_inicio': '', 'fecha_fin': ''},
                 'totales': dict(total_entradas_cantidad=Decimal('1.25'), total_entradas_monto=Decimal('12.1'),
                     total_salidas_cantidad=Decimal('1.2345'), total_salidas_monto=Decimal('2'),
                     saldo_final_cantidad=Decimal('1.000125'), saldo_final_monto=Decimal('30.1'))}
        with patch('almacen_app.views.generar_datos_libro_mayor', return_value=datos):
            response = self.client.get(reverse('almacen:libro_mayor_excel'))
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        for cell, value, format_ in [('E8',1250,'#,##0'), ('G8',1.25,'#,##0.##'),
                                      ('I8',1.2345,'#,##0.####'), ('K8',1.000125,'#,##0.######')]:
            self.assertEqual((sheet[cell].value,sheet[cell].data_type,sheet[cell].number_format),(value,'n',format_))
        for cell in ['F8','H8','J8','L8','M8']:
            self.assertEqual(sheet[cell].number_format, '"Q"#,##0.00')
